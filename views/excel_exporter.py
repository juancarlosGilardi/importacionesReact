# utils/excel_exporter.py
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime


class ExcelExporter:
    def __init__(self, db):
        self.db = db
    
    def exportar_reporte_general(self, fecha_inicio=None, fecha_fin=None):
        """Exporta un reporte general a Excel"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reporte General"
        
        # Estilos
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        center_alignment = Alignment(horizontal="center", vertical="center")
        
        # Encabezados
        headers = ["PO", "Proveedor", "Fecha Orden", "Estado", "FOB", "Gastos", "Total", "Días"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_alignment
        
        # Obtener datos
        query = """
            SELECT po.po_number, s.business_name, po.order_date, po.status,
                   po.total_fob, po.total_import_cost,
                   DATEDIFF(CURDATE(), po.order_date) as dias
            FROM purchase_orders po
            LEFT JOIN suppliers s ON po.supplier_id = s.id
            WHERE po.order_date BETWEEN %s AND %s
            ORDER BY po.order_date DESC
        """
        
        fecha_inicio = fecha_inicio or (datetime.now() - timedelta(days=30)).strftime("%Y-%m-%d")
        fecha_fin = fecha_fin or datetime.now().strftime("%Y-%m-%d")
        
        result = self.db.execute_query(query, (fecha_inicio, fecha_fin))
        
        # Llenar datos
        row = 2
        for orden in result:
            ws.cell(row=row, column=1, value=orden['po_number'])
            ws.cell(row=row, column=2, value=orden['business_name'])
            ws.cell(row=row, column=3, value=orden['order_date'])
            ws.cell(row=row, column=4, value=orden['status'])
            ws.cell(row=row, column=5, value=float(orden['total_fob'] or 0))
            ws.cell(row=row, column=6, value=float(orden['total_import_cost'] or 0) - float(orden['total_fob'] or 0))
            ws.cell(row=row, column=7, value=float(orden['total_import_cost'] or 0))
            ws.cell(row=row, column=8, value=orden['dias'])
            row += 1
        
        # Ajustar ancho de columnas
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Guardar archivo
        filename = f"reporte_general_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        wb.save(filename)
        
        return filename